# -*- coding: utf-8 -*-
"""把报告 Markdown 转成 docx（支持标题/表格/图片/代码块/引用/列表/加粗）"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cn_font(doc):
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_runs(par, text):
    """处理 **加粗** 与 `代码`。"""
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for tk in tokens:
        if not tk:
            continue
        if tk.startswith("**") and tk.endswith("**"):
            r = par.add_run(tk[2:-2])
            r.bold = True
        elif tk.startswith("`") and tk.endswith("`"):
            r = par.add_run(tk[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        else:
            par.add_run(tk)


def build(md_path, out_path, base_dir):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    set_cn_font(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        # 代码块
        if line.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            continue

        # 图片
        m = re.match(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if m:
            alt, src = m.group(1), m.group(2)
            path = os.path.join(base_dir, src.replace("/", os.sep))
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(6.4))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 表格
        if line.strip().startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for b in block:
                if re.match(r"^\|[\s\-:|]+\|$", b):
                    continue
                cells = [c.strip() for c in b.strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, rws in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        add_runs(p, rws[ci] if ci < len(rws) else "")
                        for run in p.runs:
                            run.font.size = Pt(9)
                            if ri == 0:
                                run.bold = True
                        if ri == 0:
                            from docx.oxml import OxmlElement
                            sh = OxmlElement("w:shd")
                            sh.set(qn("w:fill"), "EEEEEE")
                            cell._tc.get_or_add_tcPr().append(sh)
                doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading(level=min(lvl, 4))
            add_runs(h, m.group(2))
            i += 1
            continue

        # 引用
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            add_runs(p, line.lstrip("> ").strip())
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.italic = True
            i += 1
            continue

        # 分隔线
        if line.strip() == "---":
            doc.add_paragraph("_" * 68).alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue

        # 有序列表（手动编号，避免跨段落续号）
        m = re.match(r"^(\d+)\.\s+(.*)", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, f"{m.group(1)}. {m.group(2)}")
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.0)
        p.paragraph_format.space_after = Pt(4)
        add_runs(p, line.strip())
        i += 1

    doc.save(out_path)
    print("saved:", out_path)


if __name__ == "__main__":
    md = sys.argv[1]
    out = sys.argv[2]
    base = sys.argv[3] if len(sys.argv) > 3 else os.path.dirname(md)
    build(md, out, base)
